import os
import openpyxl
from django.core.management.base import BaseCommand
from apps.teams.models import Organisation, Department, Team, Repository, TeamDependency

# path to the excel file - it's in the project root
EXCEL_PATH = os.path.join(
    os.path.dirname(__file__),  # commands/
    '..', '..', '..', '..',  # back to project root (commands > management > teams > apps > root)
    'Agile Project Module UofW - Team Registry.xlsx'
)


class Command(BaseCommand):
    help = 'loads team data from the excel sheet into the db'

    def handle(self, *args, **kwargs):
        path = os.path.abspath(EXCEL_PATH)

        if not os.path.exists(path):
            self.stdout.write(self.style.ERROR(f"cant find excel file at: {path}"))
            return

        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.active

        # grab headers from row 1 so we can reference columns by name
        headers = [cell.value for cell in ws[1]]

        def col(row, name):
            # helper - get a cell value by column header name, strip whitespace
            idx = headers.index(name)
            val = row[idx].value
            return str(val).strip() if val else None

        # one org to rule them all - everything links back here
        org, _ = Organisation.objects.get_or_create(org_name="Sky Engineering")
        self.stdout.write("created org: Sky Engineering")

        # keep track of teams by name so we can wire up dependencies after
        team_lookup = {}

        for i, row in enumerate(ws.iter_rows(min_row=2), start=2):
            team_name = col(row, 'Team Name')

            # skip empty rows
            if not team_name:
                continue

            dept_name = col(row, 'Department') or 'Unknown'

            # get or create the department under our org
            dept, _ = Department.objects.get_or_create(
                department_name=dept_name,
                organisation=org,
            )

            # create the team - all_blank=True fields are fine to leave empty
            team, created = Team.objects.get_or_create(
                team_name=team_name,
                defaults={
                    'department': dept,
                    'skills': col(row, 'Key Skills & Technologies'),
                    'dev_focus': col(row, 'Development Focus Areas'),
                    'jira_project_name': col(row, 'Jira Project Name'),
                    'jira_board_url': _safe_url(col(row, 'Jira board Link')),
                    'slack_channel': col(row, 'Slack Channels'),
                    'standup_link': _safe_url(col(row, 'Daily Standup Time and Link')),
                    'wiki_url': _safe_url(col(row, 'Team Wiki')),
                }
            )

            # if the team already existed from a previous load, update the dept at least
            if not created:
                team.department = dept
                team.save()

            team_lookup[team_name] = team

            # repo - grab the github link if there is one
            repo_url = col(row, 'Project (codebase) (Github Repo)')
            if repo_url and repo_url.startswith('http'):
                Repository.objects.get_or_create(
                    team=team,
                    repo_url=repo_url,
                    defaults={'repo_name': team_name}
                )

            if created:
                self.stdout.write(f"  row {i}: created team '{team_name}' in dept '{dept_name}'")

        # now wire up dependencies - do this after all teams are created
        # so we can look them up by name
        self.stdout.write("\nwiring up dependencies...")
        dep_count = 0

        for row in ws.iter_rows(min_row=2):
            team_name = col(row, 'Team Name')
            if not team_name or team_name not in team_lookup:
                continue

            downstream_raw = col(row, 'Downstream Dependencies')
            dep_type = col(row, 'Dependency Type') or 'Unknown'

            if not downstream_raw:
                continue

            # dependencies can be comma separated e.g. "Team A, Team B"
            for dep_name in downstream_raw.split(','):
                dep_name = dep_name.strip()
                if dep_name in team_lookup:
                    TeamDependency.objects.get_or_create(
                        upstream_team=team_lookup[team_name],
                        downstream_team=team_lookup[dep_name],
                        defaults={'dependency_type': dep_type}
                    )
                    dep_count += 1

        self.stdout.write(self.style.SUCCESS(
            f"\ndone! loaded {len(team_lookup)} teams, {dep_count} dependencies"
        ))


def _safe_url(val):
    # only save it if it actually looks like a url, otherwise leave it blank
    if val and ('http' in val or 'www' in val):
        return val[:200]  # URLField max
    return None
